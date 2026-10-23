"""Gabarits de données de l'atelier (docs/ATELIER-CADRAGE.md, données D-1 à D-11 de docs/DECISIONS.md).

Un classeur, un onglet par donnée, colonnes alignées sur ce que Crucible vérifie à la saisie (src/domain/*.ts : checkResource,
checkExclusionRule, checkTestType, checkSubstitution, checkProject, checkMilestone, checkNeed, checkPerson, checkAgreement,
checkInspection) ; l'onglet D-8 suit exactement le format de l'import de l'historique (src/domain/history-import.ts), pour être
collé tel quel dans Planning des moyens → Importer l'historique.

    python3 scripts/atelier/gabarits.py   → docs/atelier/gabarits-donnees-crucible.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[2] / 'docs' / 'atelier' / 'gabarits-donnees-crucible.xlsx'
ROWS = 200  # lignes à remplir par onglet
FIRST = 6   # première ligne à remplir (1 titre, 2 consigne, 3 qui / où, 4 en-têtes, 5 exemple)

FONT = 'Arial'
TITLE = Font(name=FONT, size=14, bold=True, color='1F2340')
NOTE = Font(name=FONT, size=9, italic=False, color='4A4F6A')
HEAD = Font(name=FONT, size=10, bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='4B4380')
EXAMPLE = Font(name=FONT, size=10, italic=True, color='7A7F96')
EXAMPLE_FILL = PatternFill('solid', fgColor='EEEEF4')
INPUT_FILL = PatternFill('solid', fgColor='FFF7CC')
BODY = Font(name=FONT, size=10)
THIN = Side(style='thin', color='D5D6E0')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')

LISTS = {
    'type_moyen': ['zone', 'banc', 'machine', 'salle', 'équipement'],
    'granularite': ['demi-journée', 'heure'],
    'oui_non': ['oui', 'non'],
    'motif_regle': ['sécurité', 'personnel commun', 'énergie', 'nuisance', 'autre'],
    'classe': ['P1', 'P2', 'P3'],
    'criticite': ['contractuel', 'réglementaire', 'interne'],
    'niveau': ['prévision', 'demande'],
    'jour': ['lundi', 'mardi', 'mercredi', 'jeudi', 'vendredi', 'samedi', 'dimanche'],
    'indispo': ['maintenance', 'contrôle', 'panne', 'travaux', 'autre'],
}

# (nom de l'onglet, donnée, titre, consigne, qui remplit · où cela va dans Crucible, colonnes, exemple)
# Une colonne : (en-tête, largeur, liste déroulante ou None, commentaire ou None)
SHEETS = [
    ('D-1a Installations', 'D-1', 'Installations', "Une ligne par installation (laboratoire 67, PDE…). Le PDE est une installation unique découpée en zones (P-2) : ses zones sont des moyens (onglet D-1b).",
     "Direction R&D · Administration → Entités, sites, installations",
     [('Code installation', 16, None, '2 à 24 lettres, chiffres ou tirets, en majuscules (ex. LAB67, PDE).'),
      ('Nom', 28, None, None), ('Site', 18, None, 'Site du référentiel Matrix (Lyon, Le Havre, Marseille…) ; à préciser pour le laboratoire 67 (M-6).'),
      ('Entité exploitante', 18, None, 'Entité interne qui exploite l’installation (EPC2i…).'),
      ('Entités partenaires', 20, None, 'Entités extérieures qui l’utilisent aussi (EPC France…), séparées par des virgules.'),
      ('Responsable(s) (nom, e-mail)', 32, None, None), ('Coût d’acquisition (k€)', 14, None, 'Facultatif : sert à la valeur de la capacité inoccupée (P-8).'),
      ('Description', 40, None, None)],
     ['PDE', 'PDE — bâtiment pyrotechnique', 'Le Havre', 'EPC2I', '', 'Nom Prénom, prenom.nom@epc2i.fr', '3000', 'Atelier mécanique, pilote, pièces laboratoire : une zone à la fois']),
    ('D-1b Moyens', 'D-1', 'Moyens réservables', "Une ligne par moyen réservable : zone, banc, machine, salle ou équipement fixe. Les temps de préparation et de remise en état bloquent le moyen (R-5).",
     "Responsables d’installation · Fiche d’installation → Moyens et zones",
     [('Code installation', 16, None, 'Code de l’onglet D-1a.'), ('Code moyen', 16, None, '2 à 24 lettres, chiffres ou tirets (ex. PDE-PIL).'),
      ('Nom', 28, None, None), ('Type', 14, 'type_moyen', None), ('Granularité', 14, 'granularite', 'Demi-journée par défaut (F-1) ; l’heure pour les moyens qui le demandent.'),
      ('Capacité simultanée', 12, None, 'Nombre d’essais en même temps, 1 à 20 (1 par défaut, R-1).'),
      ('Préparation (min)', 12, None, '0 à 1440.'), ('Remise en état (min)', 12, None, '0 à 1440.'), ('Description', 40, None, None)],
     ['PDE', 'PDE-PIL', 'Pilote', 'zone', 'demi-journée', '1', '60', '60', 'Zone de tir du PDE']),
    ('D-1c Horaires', 'D-1', 'Horaires d’ouverture et fermetures', "Horaires par jour de la semaine (7 h–17 h par défaut, R-9) et fermetures annuelles. Laisser l’installation vide pour une fermeture de tout le site.",
     "Responsables d’installation · Fiche d’installation → Calendrier et indisponibilités",
     [('Code installation', 16, None, None), ('Jour', 12, 'jour', None), ('Ouverture (HH:MM)', 14, None, None), ('Fermeture (HH:MM)', 14, None, None),
      ('', 4, None, None), ('Fermeture : installation', 18, None, 'Vide = tout le site.'), ('Fermeture : du (AAAA-MM-JJ)', 18, None, None),
      ('Fermeture : au (AAAA-MM-JJ)', 18, None, None), ('Fermeture : motif', 28, None, None)],
     ['PDE', 'lundi', '07:00', '17:00', '', '', '2026-12-24', '2027-01-01', 'Fermeture de fin d’année']),
    ('D-2 Règles', 'D-2', 'Règles d’exclusion et de cohabitation', "Quels moyens ne peuvent pas fonctionner ensemble, et pourquoi. Chaque règle est validée par le référent sécurité (R-3) avec la référence Codex de l’analyse de risque.",
     "Référent sécurité · Règles → Règles d’exclusion",
     [('Code règle', 16, None, '2 à 24 lettres, chiffres ou tirets.'), ('Libellé', 30, None, None),
      ('Moyens concernés', 28, None, 'Au moins deux codes de moyens (onglet D-1b), séparés par des virgules.'),
      ('Maximum en même temps', 12, None, 'Combien de ces moyens peuvent fonctionner en même temps (moins que le nombre de moyens).'),
      ('Motif', 16, 'motif_regle', None), ('Pourquoi (explication)', 40, None, 'Au moins quelques mots : matière, distances, personnel commun…'),
      ('Activités concernées', 28, None, 'Toutes, ou seulement certaines (tir, pyrotechnie, présence humaine…) : sert à affiner la règle (R-2).'),
      ('Référence Codex (analyse de risque)', 22, None, None), ('Validée par (référent sécurité)', 24, None, None)],
     ['PDE-ONE-ZONE', 'PDE : une zone à la fois', 'PDE-ATM, PDE-PIL, PDE-LAB', '1', 'sécurité', 'Pas de coactivité pyrotechnique dans le bâtiment', 'Toutes', 'AR-PDE-001', 'Nom Prénom']),
    ('D-3 Types d’essai', 'D-3', 'Catalogue des types d’essai', "Les essais pratiqués : moyen nominal, durée type, compétences et équipements mobiles requis.",
     "Coordinateur · Règles → Types d’essai, compétences et équipements",
     [('Code type d’essai', 16, None, '2 à 24 lettres, chiffres ou tirets (ex. FIRING).'), ('Nom', 30, None, None),
      ('Moyen nominal', 16, None, 'Code de l’onglet D-1b.'), ('Durée type (demi-journées)', 14, None, '1 à 400.'),
      ('Compétences requises', 30, None, 'Code de la compétence × nombre de personnes, séparés par des virgules (ex. PYRO×2, INSTR×1). Codes de l’onglet D-6b.'),
      ('Équipements mobiles requis', 24, None, 'Codes de l’onglet D-6d, séparés par des virgules.'), ('Description', 40, None, None)],
     ['FIRING', 'Tir de qualification', 'PDE-PIL', '2', 'PYRO×2', 'CAM-01', 'Tir instrumenté, caméra rapide']),
    ('D-4 Substitutions', 'D-4', 'Substitutions possibles', "Quels essais peuvent se faire sur un autre moyen, à quelles conditions. Validées par le responsable de l’installation d’accueil (et le référent sécurité si matières dangereuses, R-4, R-11).",
     "Coordinateur, responsables d’installation · Règles → Substitutions",
     [('Code type d’essai', 16, None, 'Onglet D-3.'), ('Moyen substitut', 16, None, 'Code de l’onglet D-1b (autre que le moyen nominal).'),
      ('Conditions', 40, None, 'Au moins quelques mots.'), ('Référence Codex de la justification', 22, None, None),
      ('Référent technique', 22, None, None), ('Matières dangereuses', 12, 'oui_non', None), ('Validée par (installation d’accueil)', 24, None, None)],
     ['FORMULATION', 'LAB67-ANA', 'Échantillons < 50 g, hotte dédiée', 'JT-LAB-014', 'Nom Prénom', 'non', 'Nom Prénom']),
    ('D-5a Projets', 'D-5', 'Projets', "Les projets R&D du pilote. La classe de priorité est fixée par la direction R&D (A-5) ; un projet confidentiel n’est visible que de son équipe et de la direction R&D (U-3).",
     "Direction R&D · Administration → Projets",
     [('Code projet', 16, None, '2 à 24 lettres, chiffres ou tirets.'), ('Nom', 28, None, None), ('Entité', 14, None, 'EPC2I, EPCFR…'),
      ('Classe de priorité', 12, 'classe', 'P1 stratégique, P2 standard, P3 exploratoire.'), ('Confidentiel', 12, 'oui_non', None),
      ('Chef de projet (nom, e-mail)', 30, None, None), ('Code projet Keystone', 16, None, 'Facultatif.'), ('Code programme Codex', 16, None, 'Facultatif.')],
     ['RD-ORION', 'Orion', 'EPC2I', 'P1', 'non', 'Nom Prénom, prenom.nom@epc2i.fr', '', '']),
    ('D-5b Jalons', 'D-5', 'Jalons des projets', "Les jalons qui portent des essais. Criticité : contractuel (engagement client daté), réglementaire, interne.",
     "Chefs de projet · Fiche projet → Jalons",
     [('Code projet', 16, None, 'Onglet D-5a.'), ('Jalon', 30, None, None), ('Date (AAAA-MM-JJ)', 16, None, None), ('Criticité', 14, 'criticite', None)],
     ['RD-ORION', 'Revue de qualification', '2027-01-15', 'interne']),
    ('D-5c Besoins', 'D-5', 'Besoins d’essais à 6–12 mois', "Les essais prévus, même approximatifs (« ~15 jours de PDE avant mars ») : une prévision. Une demande est un essai précis à planifier, au moins 10 jours ouvrés avant (F-5).",
     "Chefs de projet · Mes besoins d’essais",
     [('Code projet', 16, None, 'Onglet D-5a.'), ('Niveau', 12, 'niveau', None), ('Essais nécessaires', 34, None, '3 à 200 caractères.'),
      ('Type d’essai', 16, None, 'Onglet D-3, facultatif.'), ('Moyen', 16, None, 'Onglet D-1b.'), ('Demi-journées', 12, None, '1 à 400.'),
      ('Jalon servi', 24, None, 'Onglet D-5b, facultatif : sans échéance, c’est la date du jalon.'),
      ('Au plus tôt (AAAA-MM-JJ)', 16, None, None), ('Échéance (AAAA-MM-JJ)', 16, None, None), ('Référence Codex du protocole', 20, None, None), ('Note', 30, None, None)],
     ['RD-ORION', 'prévision', 'Campagne de tirs avant la revue', 'FIRING', 'PDE-PIL', '6', 'Revue de qualification', '2027-01-04', '', '', '']),
    ('D-6a Personnes', 'D-6', 'Personnes qui réalisent les essais', "Opérateurs, techniciens, pyrotechniciens : entité, matricule (pour relier Cursus), installations où ils travaillent.",
     "Responsables d’installation · Personnes et compétences",
     [('Nom', 26, None, None), ('Entité', 14, None, None), ('Matricule', 14, None, 'Facultatif ; relie la personne à Cursus.'),
      ('Installations', 22, None, 'Codes de l’onglet D-1a, séparés par des virgules.'), ('Contraintes horaires', 30, None, None)],
     ['Nom Prénom', 'EPC2I', '000123', 'PDE', '']),
    ('D-6b Compétences', 'D-6', 'Compétences et habilitations', "Les compétences qu’un type d’essai peut demander (habilitation pyrotechnique, instrumentation…).",
     "Coordinateur · Personnes et compétences → Compétences",
     [('Code compétence', 16, None, '2 à 24 lettres, chiffres ou tirets.'), ('Nom', 34, None, None), ('Titre Cursus (identifiant)', 22, None, 'Facultatif.')],
     ['PYRO', 'Habilitation pyrotechnique', '']),
    ('D-6c Détenues', 'D-6', 'Compétences détenues', "Qui détient quelle compétence, jusqu’à quand.",
     "Responsables d’installation · Personnes et compétences",
     [('Personne (nom ou matricule)', 26, None, 'Onglet D-6a.'), ('Code compétence', 16, None, 'Onglet D-6b.'), ('Valide jusqu’au (AAAA-MM-JJ)', 18, None, 'Vide = sans limite.')],
     ['000123', 'PYRO', '2027-06-30']),
    ('D-6d Équipements', 'D-6', 'Équipements mobiles partagés', "Instrumentation, caméras, chaînes d’acquisition déplaçables entre moyens (R-8).",
     "Responsables d’installation · Fiche d’installation → Équipements mobiles",
     [('Code équipement', 16, None, '2 à 24 lettres, chiffres ou tirets.'), ('Nom', 30, None, None), ('Installation', 16, None, 'Onglet D-1a.')],
     ['CAM-01', 'Caméra rapide', 'PDE']),
    ('D-7 Accord EPC France', 'D-7', 'Accord de partage', "S’il existe : quota de jours par mois, délai de prévenance, priorités, refacturation. Sans accord écrit, l’écrire sur la ligne d’exemple (A-3).",
     "Direction R&D · Règles → Accords",
     [('Installation', 16, None, None), ('Entité partenaire', 16, None, None), ('Quota (jours / mois)', 14, None, '0 à 31, vide = sans quota.'),
      ('Délai de prévenance (jours ouvrés)', 16, None, '0 à 60, 10 par défaut (F-5).'), ('Résumé de l’accord', 50, None, None), ('Contact', 24, None, None)],
     ['LAB67', 'EPCFR', '8', '10', 'Aucun accord écrit à ce jour', 'Nom Prénom']),
    ('D-8 Historique', 'D-8', 'Historique des réservations (6 à 12 mois)', "Format de l’import de Crucible : copier les lignes remplies (sans la ligne d’exemple) et les coller dans Planning des moyens → Importer l’historique. Dates et heures « AAAA-MM-JJ HH:MM ». Heures réelles vides : l’essai a eu lieu comme prévu ; « absent » dans le début réel : non tenu.",
     "Coordinateur (les responsables d’installation fournissent les lignes) · Planning des moyens → Importer l’historique",
     [('resource · code du moyen', 18, None, 'Code de l’onglet D-1b. (L’en-tête commence par « resource » : l’import l’ignore s’il est collé.)'),
      ('Code projet', 14, None, 'Onglet D-5a.'), ('Début prévu', 18, None, 'AAAA-MM-JJ HH:MM'), ('Fin prévue', 18, None, 'AAAA-MM-JJ HH:MM'),
      ('Début réel', 18, None, 'AAAA-MM-JJ HH:MM, vide, ou « absent ».'), ('Fin réelle', 18, None, 'AAAA-MM-JJ HH:MM ou vide.'), ('Objet de l’essai', 36, None, '3 caractères au moins.')],
     ['PDE-PIL', 'RD-ORION', '2026-09-14 07:00', '2026-09-14 17:00', '2026-09-14 07:40', '2026-09-14 15:30', 'Tirs, série 1']),
    ('D-9 Chiffre des 30 %', 'D-9', 'Origine du chiffre « PDE occupé à 30 % »', "Qui l’a calculé, sur quelle période, avec quelle formule (P-5).",
     "Responsable produit · Questionnaire, bloc D-9",
     [('Auteur', 24, None, None), ('Période', 24, None, None), ('Formule', 40, None, None), ('Source des données', 30, None, None), ('Remarques', 40, None, None)],
     ['Nom Prénom', '2025', 'heures d’essai / heures ouvrées', 'Agenda Outlook du PDE', '']),
    ('D-10 Contrôles', 'D-10', 'Contrôles réglementaires et maintenances', "Contrôles périodiques des moyens (un contrôle échu bloque le moyen, R-6) et maintenances prévues (indisponibilités).",
     "Responsables d’installation · Fiche d’installation → Contrôles périodiques",
     [('Code moyen', 16, None, 'Onglet D-1b.'), ('Contrôle', 32, None, None), ('Périodicité (mois)', 12, None, '1 à 120.'),
      ('Dernier contrôle (AAAA-MM-JJ)', 18, None, 'Pas dans le futur.'), ('Référence Codex', 18, None, None), ('', 4, None, None),
      ('Maintenance : type', 14, 'indispo', None), ('Maintenance : libellé', 26, None, None), ('Du (AAAA-MM-JJ)', 16, None, None), ('Au (AAAA-MM-JJ)', 16, None, None)],
     ['PDE-PIL', 'Vérification périodique des installations électriques', '12', '2025-11-16', 'CTRL-ELEC-PDE', '', 'maintenance', 'Révision annuelle', '2027-02-08', '2027-02-09']),
    ('D-11 Participants', 'D-11', 'Participants à l’atelier', "Qui inviter (docs/ATELIER-CADRAGE.md) : direction R&D, responsables d’installation, référent sécurité, chefs de projet, représentant d’EPC France, coordinateur, DSI.",
     "Responsable produit · Questionnaire, bloc D-11",
     [('Nom', 26, None, None), ('Fonction', 30, None, None), ('Entité', 14, None, None), ('E-mail', 30, None, None), ('Présent', 10, 'oui_non', None)],
     ['Nom Prénom', 'Référent sécurité pyrotechnique', 'EPC2I', 'prenom.nom@epc2i.fr', 'oui']),
]


def style_range(ws, row, ncols, font, fill=None, border=True):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = font
        cell.alignment = WRAP
        if fill is not None:
            cell.fill = fill
        if border:
            cell.border = BOX


def build():
    wb = Workbook()
    readme = wb.active
    readme.title = 'Lisez-moi'

    lists = wb.create_sheet('Listes')
    for i, (name, values) in enumerate(LISTS.items(), start=1):
        lists.cell(row=1, column=i, value=name).font = Font(name=FONT, bold=True)
        for j, v in enumerate(values, start=2):
            lists.cell(row=j, column=i, value=v).font = BODY
    lists.sheet_state = 'hidden'
    list_ref = {name: f"Listes!${get_column_letter(i)}$2:${get_column_letter(i)}${len(v) + 1}" for i, (name, v) in enumerate(LISTS.items(), start=1)}

    for name, data, title, note, where, cols, example in SHEETS:
        ws = wb.create_sheet(name)
        n = len(cols)
        ws['A1'] = f'{data} · {title}'
        ws['A1'].font = TITLE
        ws['A2'] = note
        ws['A2'].font = NOTE
        ws['A2'].alignment = WRAP
        ws['A3'] = f'Qui remplit · où cela va dans Crucible : {where}'
        ws['A3'].font = Font(name=FONT, size=9, bold=True, color='4B4380')
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(n, 4))
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(n, 4))
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max(n, 4))
        ws.row_dimensions[2].height = 42
        for c, (head, width, lst, comment) in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(c)].width = width
            cell = ws.cell(row=4, column=c, value=head or None)
            if head:
                cell.font = HEAD
                cell.fill = HEAD_FILL
                cell.alignment = Alignment(wrap_text=True, vertical='center')
                cell.border = BOX
                if comment:
                    cell.comment = Comment(comment, 'Crucible')
            ex = ws.cell(row=5, column=c, value=(example[c - 1] or None) if head else None)
            if head:
                ex.font = EXAMPLE
                ex.fill = EXAMPLE_FILL
                ex.border = BOX
                ex.number_format = '@'
            for r in range(FIRST, FIRST + ROWS):
                cell = ws.cell(row=r, column=c)
                if head:
                    cell.fill = INPUT_FILL
                    cell.border = BOX
                    cell.font = BODY
                    cell.number_format = '@'  # codes, dates et heures gardés tels quels (pas de conversion par le tableur)
            if lst:
                dv = DataValidation(type='list', formula1=list_ref[lst], allow_blank=True, showErrorMessage=True,
                                    errorTitle='Valeur', error='Choisissez une valeur de la liste.')
                ws.add_data_validation(dv)
                dv.add(f'{get_column_letter(c)}{FIRST}:{get_column_letter(c)}{FIRST + ROWS - 1}')
        ws.row_dimensions[4].height = 34
        ws.cell(row=5, column=n + 1, value='← exemple, à ne pas recopier').font = EXAMPLE
        ws.freeze_panes = f'A{FIRST}'

    # Lisez-moi : objet, conventions, suivi du remplissage (formules).
    r = readme
    r.column_dimensions['A'].width = 24
    r.column_dimensions['B'].width = 10
    r.column_dimensions['C'].width = 46
    r.column_dimensions['D'].width = 54
    r.column_dimensions['E'].width = 14
    r['A1'] = 'Crucible · gabarits de données de l’atelier'
    r['A1'].font = TITLE
    lines = [
        'Ce classeur recueille les données D-1 à D-11 nécessaires au pilote de Crucible (laboratoire 67, PDE et ses zones). Un onglet par donnée ; chaque colonne correspond à un champ que Crucible vérifie à la saisie.',
        'Remplir les cellules jaunes, une ligne par élément. La ligne grise de chaque onglet est un exemple : ne pas la recopier. Les en-têtes portent une note (survoler la cellule) qui précise le format attendu.',
        'Conventions : codes en majuscules, 2 à 24 lettres, chiffres ou tirets (PDE, PDE-PIL, RD-ORION) ; dates AAAA-MM-JJ ; heures HH:MM ; listes séparées par des virgules. Les cellules sont au format texte pour que le tableur ne transforme ni les codes ni les dates.',
        'Les codes relient les onglets : un moyen (D-1b) appartient à une installation (D-1a), un type d’essai (D-3) a un moyen nominal, un besoin (D-5c) cite un projet, un moyen et éventuellement un jalon.',
        'L’onglet D-8 (historique) suit exactement le format de l’import de Crucible : ses lignes se collent telles quelles dans Planning des moyens → Importer l’historique. C’est la base de la mesure de référence du pilote (X-2).',
        'À renvoyer au responsable produit ; les réponses peuvent aussi être collées dans les blocs « Données » du questionnaire « Décisions Crucible ».',
    ]
    for i, text in enumerate(lines, start=3):
        r.cell(row=i, column=1, value=text).font = BODY
        r.cell(row=i, column=1).alignment = WRAP
        r.merge_cells(start_row=i, start_column=1, end_row=i, end_column=5)
        r.row_dimensions[i].height = 44
    head_row = 3 + len(lines) + 1
    for c, h in enumerate(['Onglet', 'Donnée', 'Contenu', 'Qui remplit · où cela va dans Crucible', 'Lignes remplies'], start=1):
        cell = r.cell(row=head_row, column=c, value=h)
        cell.font = HEAD
        cell.fill = HEAD_FILL
        cell.border = BOX
    for i, (name, data, title, _note, where, _cols, _ex) in enumerate(SHEETS, start=head_row + 1):
        r.cell(row=i, column=1, value=name)
        r.cell(row=i, column=2, value=data)
        r.cell(row=i, column=3, value=title)
        r.cell(row=i, column=4, value=where)
        # Lignes remplies : la première colonne, de la première ligne à remplir à la dernière (l’exemple n’est pas compté).
        r.cell(row=i, column=5, value=f"=COUNTA('{name}'!A{FIRST}:A{FIRST + ROWS - 1})")
        style_range(r, i, 5, BODY)
    legend = head_row + len(SHEETS) + 2
    r.cell(row=legend, column=1, value='Légende').font = Font(name=FONT, bold=True)
    r.cell(row=legend + 1, column=1, value='À remplir').fill = INPUT_FILL
    r.cell(row=legend + 2, column=1, value='Exemple').fill = EXAMPLE_FILL
    for k in (1, 2):
        r.cell(row=legend + k, column=1).font = BODY
        r.cell(row=legend + k, column=1).border = BOX
    r.freeze_panes = None

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == '__main__':
    build()
